"""Builds lista-importerow.xlsx and lista-importerow.md from the lists below.

Run: python3 build.py   (requires openpyxl)
"""
from pathlib import Path
from urllib.parse import urlsplit

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

P = True  # Polish producer

# Each company appears in exactly one category, and only Polish(-registered) companies are listed.
LISTS = [
    ("Bielizna", "Największy importer bielizny", [
        ("LPP S.A. (Sinsay, Reserved, Cropp, House, Mohito)", "https://www.lpp.com"),
        ("Pepco Poland", "https://pepco.pl"),
        ("Biedronka (Jeronimo Martins Polska)", "https://www.biedronka.pl"),
        ("Esotiq & Henderson S.A.", "https://henderson.pl"),
        ("Atlantic", "https://www.atlantic.pl", P),
        ("Moraj", "https://moraj.pl"),
        ("4F (OTCF S.A.)", "https://4f.com.pl"),
        ("Modivo", "https://modivo.pl"),
        ("Cornette", "https://cornette.pl", P),
        ("Gatta", "https://gatta.pl", P),
        ("Julimex", "https://julimex.pl", P),
        ("Redo Polska Sp. z o.o.", "https://redopolska.pl"),
        ("Nattalex – hurtownia bielizny (Chorzów)", "https://nattalex.pl"),
        ("Noviti Sp. z o.o.", "https://noviti.pl", P),
        ("LA/MA (Lama)", "https://lama.com.pl", P),
        ("Babell", "https://babell.com.pl", P),
        ("Wol-Bar", "https://wol-bar.com.pl", P),
        ("Brubeck", "https://brubeck.pl", P),
        ("BieliznaCentrum.pl", "https://bieliznacentrum.pl"),
        ("Mada.pl", "https://www.mada.pl"),
        ("Kentex", "https://kentex.com.pl"),
        ("Hurtownia Bielizny i Rajstop Oleńka (Marki)", "https://www.hurtownia-olenka.pl"),
        ("Hurtownia TEA (Łódź)", "https://tea.com.pl"),
        ("Anstel (Częstochowa)", "https://www.anstel.pl"),
        ("De Lafense (dawniej Forex Lingerie)", "https://www.delafense.com", P),
        ("Szata", "https://szata.pl", P),
        ("Martex", "https://martexsklep.pl", P),
        ("Gucio", "https://gucio.com.pl", P),
    ]),
    ("Skarpety", "Największy importer skarpet", [
        ("HalfPrice (Grupa CCC)", "https://halfprice.eu/pl"),
        ("JS Sp. z o.o. (Lublin)", "https://www.js.com.pl"),
        ("Smyk", "https://www.smyk.com/pl/pl"),
        ("Soxo", "https://soxo.pl"),
        ("Skarpex", "https://www.skarpex.pl"),
        ("All-Best", "https://www.allbest.pl"),
        ("Rainbow Socks", "https://rainbowsocks.pl"),
        ("Krebo Sp. z o.o.", "https://krebo.com.pl", P),
        ("Wola", "https://wola.pl", P),
        ("Steven", "https://steven.pl", P),
        ("Marilyn", "https://e-marilyn.pl", P),
        ("Zooksy", "https://zooksysocks.com", P),
        ("Intenso / TAK (Łowicz)", "https://sklepintenso.pl", P),
        ("Many Mornings", "https://manymornings.com", P),
        ("Spox Sox", "https://spoxsox.pl", P),
        ("Milena", "https://milena-skarpety.pl", P),
        ("Gram-Pol", "https://gram-pol.pl", P),
        ("Cerber (Łowicz)", "https://www.socks.pl", P),
        ("Todo Socks", "https://www.todosocks.pl", P),
        ("Krystjan (Crazy Socks)", "https://krystjan.com.pl", P),
        ("Socks Factory Sp. z o.o.", "https://socksfactory.pl", P),
        ("Rajan", "https://www.rajan.com.pl", P),
        ("Skarpol", "https://skarpol.com", P),
        ("Bornpol", "https://bornpol.pl", P),
        ("Eltom", "https://www.skarpety-eltom.pl", P),
        ("Bratex", "http://www.bratexskarpetki.pl", P),
    ]),
    ("Piżamy", "Największy importer piżam", [
        ("Key", "https://key.com.pl", P),
        ("Italian Fashion", "https://www.italianfashion.pl", P),
        ("Taro", "https://taro.pl", P),
        ("Mewa (Mewa Lingerie, Visti, James Bradley)", "https://mewa.com.pl", P),
        ("Hurtownia WEGA (Wrocław)", "https://www.wega.wroc.pl"),
        ("Doctor Nap (dawniej Dobranocka)", "https://doctornap.com/pl", P),
        ("Regina", "https://regina.pl", P),
        ("Kuba", "https://kubapizama.pl", P),
        ("M-Max", "https://m-max.com.pl", P),
        ("Martel", "https://martel-pphu.pl", P),
        ("VMP Group (Momenti Per Me)", "https://vmp.com.pl"),
        ("Muzzy", "https://muzzy.pl", P),
        ("Cana Nouvelle", "https://cana.pl", P),
        ("Donna", "https://www.bedonna.pl", P),
        ("Sesto Senso", "https://sestosenso24.pl", P),
        ("Coemi", "https://coemi.pl", P),
        ("Dkaren", "https://dkaren.pl", P),
        ("Kalimo", "https://kalimo.eu", P),
        ("Stefani", "https://stefani.com.pl", P),
        ("Dorota (Łódź)", "http://dorota-bielizna.pl", P),
        ("Born2rest", "https://born2rest.com", P),
        ("Lorin", "https://sklep.lorin.eu", P),
        ("Betina", "https://betinasklep.pl", P),
        ("AS Piżamy", "https://aspizamy.pl", P),
        ("Bohomoss", "https://bohomoss.com", P),
        ("Sleep Art", "https://sleepart.pl", P),
        ("Sunday Mood", "https://sundaymood.pl", P),
    ]),
    ("Biustonosze", "Największy importer staników", [
        ("Intymna.pl / Matterhorn (Grupa Modne Zakupy)", "https://www.intymna.pl"),
        ("Gorteks", "https://www.gorteks.com.pl", P),
        ("Obsessive", "https://obsessive.com"),
        ("Kontri.pl", "https://www.kontri.pl"),
        ("Teyli", "https://www.teyli.eu"),
        ("Ava Lingerie", "https://avalingerie.pl", P),
        ("Krisline", "https://krisline.com", P),
        ("Gaia", "https://www.gaia.com.pl", P),
        ("Gorsenia", "https://gorsenia.pl", P),
        ("Alles", "https://alles.pl", P),
        ("Kinga Lingerie", "https://kinga.com.pl", P),
        ("Dalia Lingerie", "https://dalia.pl", P),
        ("Mefemi (dawniej Nipplex)", "https://mefemi.pl", P),
        ("Ewa Bien", "https://ewabien.com", P),
        ("Axami", "https://axami.pl", P),
        ("Samanta", "https://samanta.pl", P),
        ("Mitex", "https://mitex.pl", P),
        ("MAT Lingerie", "https://matlingerie.pl", P),
        ("Konrad Lingerie", "https://konrad-bielizna.pl", P),
        ("Lupoline", "https://esklep.lupoline.pl", P),
        ("Corin", "https://corin.eu", P),
        ("Ewa Michalak", "https://www.ewa-michalak.pl", P),
        ("Nessa", "https://nessa.com.pl", P),
        ("Feba", "https://feba.pl", P),
        ("Comexim Lingerie", "https://comexim.pl", P),
        ("Lapinee", "https://lapinee.pl"),
        ("Undress Code", "https://pl.undress-code.com", P),
        ("God Save Queens", "https://godsavequeens.com"),
    ]),
    ("Bielizna z Turcji", "Największy importer bielizny z Turcji", [
        ("Duman Polska Sp. z o.o.", "https://www.dumanpolska.com.pl"),
        ("Donella Sp. z o.o. (Wólka Kosowska)", "https://donella.pl"),
        ("Raj-Pol (Częstochowa)", "https://raj-pol.com.pl", P),
        ("Belbest Sp. z o.o.", "https://belbest.com.pl"),
        ("Luka Hurt Sp. z o.o.", "https://hurtownialuka.pl"),
        ("Onlinehurt.pl (Tivaro Sp. z o.o.)", "https://onlinehurt.pl"),
        ("Felix Hardy", "https://felixhardy.com"),
        ("JohnFrank.pl (Bardotka, Wejherowo)", "https://johnfrank.pl"),
        ("ASG3HURT", "https://asg3hurt.com"),
        ("CentrumHurt", "https://centrumhurt.pl"),
        ("HurtowniaWolka.com / WolkaCentrum.com", "https://hurtowniawolka.com"),
        ("S&D Bielizna Turecka (GD Center)", "https://gdonline.pl/SD-Bielizna-Turecka/pr/100"),
        ("KOZA Underwear Polska (GD Center)", "https://gdonline.pl/KOZA-Underwear/pr/47"),
        ("Ka-Ma Hurtownia Tekstyliów (Andrychów)", "https://kama-andrychow.com"),
        ("Hurtownia Edgar (Wrocław)", "https://hurtowniaedgar.pl"),
        ("Larado", "http://larado.pl"),
        ("BETA&PLUS", "https://www.hurt-odziez.pl"),
        ("24Hurt", "https://24hurt.eu"),
        ("Touches", "https://touches.pl"),
        ("nbielizna.pl (P.H. Look)", "https://nbielizna.pl"),
        ("Modny Sklep (Wrocław)", "https://www.modny-sklep.pl"),
        ("Biuścik – Bielizna dla Ciebie", "https://www.biuścik.pl"),
        ("Datura (Tarnawa)", "https://sklep-datura.pl"),
    ]),
    ("Piżamy z Turcji", "Największy importer piżam z Turcji", [
        ("NANA Vienetta Polska", "https://vienettahurt.pl"),
        ("Vienetta – Klaudia Cieplak", "https://vienetta.pl"),
        ("Morella (HASSEM Sp. z o.o.)", "https://morellahurt.pl"),
        ("Hurtownia Duet-2", "https://www.hurtduet2.pl"),
        ("ANIAPHU – Hurtownia Piżam", "https://www.piżamyhurt.pl"),
        ("F.H. Viki / F.H. Tomi", "https://www.posciel-tac.pl"),
        ("Tuva Home Sp. z o.o.", "https://sklep.tuvahome.pl"),
        ("HurtWolka24.pl", "https://hurtwolka24.pl"),
        ("Wolka-Centrum.eu", "https://wolka-centrum.eu"),
        ("ehurtWolka.pl", "https://ehurtwolka.pl"),
        ("e24wolka.eu", "https://e24wolka.eu"),
        ("Wolka24Hurt.eu", "https://wolka24hurt.eu"),
        ("Loomac", "https://loomac.com"),
        ("Ottomania", "https://ottomania.pl"),
        ("Jabos.pl", "https://jabos.pl"),
        ("Sklep Megi (Tarnów)", "https://sklepmegitarnow.pl"),
        ("Gabrysiowy Zakątek", "https://gabrysiowyzakatek.pl"),
        ("Takistyl", "https://takistyl.pl"),
        ("Turecki Sklep", "https://tureckisklep.pl"),
        ("Bucik Sklep", "https://buciksklep.pl"),
    ]),
]

NOTES = [
    "Każda firma występuje tylko w jednej kategorii. Na listach są wyłącznie polskie firmy (zagraniczne sieci i marki, np. H&M, Primark, KiK, Lidl, LC Waikiki, zostały usunięte).",
    "Kolejność jest szacunkowa (wielkość i znaczenie firmy) – brak publicznych danych o wolumenie importu w tych kategoriach.",
    "Wszystkie adresy WWW zostały sprawdzone (strona otwarta lub potwierdzona w wynikach wyszukiwania) w październiku 2026.",
    "Biedronka (Jeronimo Martins Polska) i Pepco mają kapitał zagraniczny, ale centrale w Polsce – do decyzji, czy zostają.",
    "Duman Polska, Donella i Tuva Home to polskie spółki tureckich producentów – mogą być partnerem, ale też konkurencją dla innych tureckich dostawców.",
    "Dwie firmy podają się za oficjalnego dystrybutora Vienetta (NANA Vienetta Polska – Ptak Rzgów; Klaudia Cieplak – EACC Wólka).",
    "S&D Bielizna Turecka i KOZA Underwear mają tylko profile na gdonline.pl (stoiska w GD Center, Wólka Kosowska).",
    "Adresy z polskimi literami: www.piżamyhurt.pl (punycode: www.xn--piamyhurt-ccc.pl) i www.biuścik.pl (punycode: www.xn--biucik-5ib.pl).",
]

HERE = Path(__file__).parent


def rows(entries):
    for i, (name, url, *flag) in enumerate(entries, 1):
        yield i, name, url, "tak" if flag else ""


def ascii_url(url):
    """Excel can't open hyperlinks with non-ASCII host names, so use punycode for those."""
    parts = urlsplit(url)
    return parts._replace(netloc=parts.netloc.encode("idna").decode()).geturl()


def build_xlsx():
    wb = Workbook()
    wb.remove(wb.active)
    head_fill = PatternFill("solid", fgColor="081F45")
    for sheet, title, entries in LISTS:
        ws = wb.create_sheet(sheet)
        ws.append([title])
        ws["A1"].font = Font(bold=True, size=13)
        ws.append(["#", "Nazwa", "WWW", "Polski producent"])
        for cell in ws[2]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = head_fill
        for i, name, url, prod in rows(entries):
            ws.append([i, name, url, prod])
            link = ws.cell(row=ws.max_row, column=3)
            link.hyperlink = ascii_url(url)
            link.font = Font(color="0D4C9A", underline="single")
        for col, width in zip("ABCD", (5, 52, 50, 17)):
            ws.column_dimensions[col].width = width
        for row in ws.iter_rows(min_row=3, min_col=4, max_col=4):
            row[0].alignment = Alignment(horizontal="center")
        ws.freeze_panes = "A3"
    wb.save(HERE / "lista-importerow.xlsx")


def build_md():
    out = ["# Importerzy bielizny, skarpet i piżam – Polska", "",
           "Kolumna „PL prod.” = polski producent. Pełna lista w `lista-importerow.xlsx` (jedna zakładka na kategorię).", ""]
    for _, title, entries in LISTS:
        out += [f"## {title} ({len(entries)})", "", "| # | Nazwa | WWW | PL prod. |", "|---|---|---|---|"]
        out += [f"| {i} | {name} | {url} | {prod} |" for i, name, url, prod in rows(entries)]
        out.append("")
    out += ["## Uwagi", ""] + [f"- {note}" for note in NOTES] + [""]
    (HERE / "lista-importerow.md").write_text("\n".join(out), encoding="utf-8")


def check_unique():
    seen = {}
    for _, title, entries in LISTS:
        for name, url, *_ in entries:
            parts = urlsplit(url)
            host = parts.netloc.removeprefix("www.") + (parts.path if parts.netloc == "gdonline.pl" else "")
            assert host not in seen, f"{name} ({host}) in both '{seen.get(host)}' and '{title}'"
            seen[host] = title


if __name__ == "__main__":
    check_unique()
    build_xlsx()
    build_md()
    for _, title, entries in LISTS:
        print(f"{len(entries):3d}  {title}")
