"""Builds the Exercise 3 report (.docx) from a content module.

Usage: python3 build.py <en|pl> <output.docx>
"""

import importlib
import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

FONT = "Times New Roman"
HERE = Path(__file__).resolve().parent


def set_font(style_or_run, name=FONT, size=None, bold=None, color=None):
    font = style_or_run.font
    font.name = name
    rpr = style_or_run.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        rfonts.attrib.pop(qn(attr), None)
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rfonts.set(qn(attr), name)
    if size is not None:
        font.size = Pt(size)
    if bold is not None:
        font.bold = bold
    if color is not None:
        font.color.rgb = color


def setup_styles(doc):
    normal = doc.styles["Normal"]
    set_font(normal, size=12)
    pf = normal.paragraph_format
    pf.line_spacing = 1.5
    pf.space_after = Pt(6)
    pf.space_before = Pt(0)
    pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    for name, size, before, after in (
        ("Heading 1", 16, 18, 8),
        ("Heading 2", 14, 12, 6),
        ("Heading 3", 12, 10, 4),
    ):
        st = doc.styles[name]
        set_font(st, size=size, bold=True, color=RGBColor(0, 0, 0))
        st.font.italic = False
        st.paragraph_format.space_before = Pt(before)
        st.paragraph_format.space_after = Pt(after)
        st.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
        st.paragraph_format.keep_with_next = True
        st.paragraph_format.line_spacing = 1.15

    for name in ("List Bullet", "List Number"):
        st = doc.styles[name]
        set_font(st, size=12)
        st.paragraph_format.space_after = Pt(3)
        st.paragraph_format.line_spacing = 1.5

    cap = doc.styles["Caption"]
    set_font(cap, size=10, bold=False, color=RGBColor(0, 0, 0))
    cap.font.italic = True
    cap.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_before = Pt(4)
    cap.paragraph_format.space_after = Pt(10)
    cap.paragraph_format.line_spacing = 1.0


def add_rich(paragraph, text, size=None):
    """Adds text supporting **bold** and [[placeholder]] (highlighted) markup."""
    for part in re.split(r"(\*\*.+?\*\*|\[\[.+?\]\])", text):
        if not part:
            continue
        if part.startswith("**"):
            run = paragraph.add_run(part[2:-2])
            run.bold = True
        elif part.startswith("[["):
            run = paragraph.add_run(part[2:-2])
            run.font.highlight_color = WD_COLOR_INDEX.YELLOW
        else:
            run = paragraph.add_run(part)
        if size:
            run.font.size = Pt(size)


def add_field(paragraph, instr):
    run = paragraph.add_run()
    for tag, text in (("begin", None), ("instr", instr), ("separate", None), ("end", None)):
        if tag == "instr":
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = text
        else:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), tag)
        run._r.append(el)
        if tag == "separate":
            t = OxmlElement("w:t")
            t.text = "1"
            run._r.append(t)


def shade(cell, hex_color):
    tcpr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tcpr.append(shd)


def row_prop(row, tag):
    trpr = row._tr.get_or_add_trPr()
    el = OxmlElement(tag)
    el.set(qn("w:val"), "true")
    trpr.append(el)


def repeat_header(row):
    row_prop(row, "w:tblHeader")


def cell_text(cell, text, bold=False, size=10, bullets=False):
    cell.paragraphs[0].text = ""
    items = text if isinstance(text, list) else [text]
    for i, item in enumerate(items):
        p = cell.paragraphs[0] if i == 0 else cell.add_paragraph()
        p.paragraph_format.line_spacing = 1.0
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
        add_rich(p, ("• " if bullets else "") + item, size=size)
        if bold:
            p.paragraph_format.keep_with_next = True
            for r in p.runs:
                r.bold = True


def add_table(doc, headers, rows, widths=None, header_fill="D9E2F3", bullets=False):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(headers):
        c = table.rows[0].cells[i]
        cell_text(c, h, bold=True)
        shade(c, header_fill)
    repeat_header(table.rows[0])
    for row in rows:
        new_row = table.add_row()
        row_prop(new_row, "w:cantSplit")
        cells = new_row.cells
        for i, val in enumerate(row):
            cell_text(cells[i], val, bullets=bullets and isinstance(val, list))
    if widths:
        for row in table.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = Cm(w)
    return table


def title_page(doc, tp):
    def line(text, size=12, bold=False, before=0, after=0, align=WD_ALIGN_PARAGRAPH.CENTER):
        p = doc.add_paragraph()
        p.alignment = align
        p.paragraph_format.space_before = Pt(before)
        p.paragraph_format.space_after = Pt(after)
        p.paragraph_format.line_spacing = 1.15
        add_rich(p, text, size=size)
        if bold:
            for r in p.runs:
                r.bold = True
        return p

    line(tp["university"], 14, bold=True, before=24)
    line(tp["faculty"], 12, after=60)
    line(tp["course"], 14, bold=True)
    line(tp["exercise"], 12, after=48)
    line(tp["title"], 20, bold=True, after=8)
    line(tp["subtitle"], 13, after=110)
    for label, value in tp["fields"]:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.left_indent = Cm(8)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.15
        r = p.add_run(label + ": ")
        r.bold = True
        add_rich(p, value)
    line(tp["place_date"], 12, before=80)
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def build(lang, out_path):
    content = importlib.import_module(f"content_{lang}")
    doc = Document()
    setup_styles(doc)

    sec = doc.sections[0]
    sec.page_height, sec.page_width = Cm(29.7), Cm(21.0)
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(sec, side, Cm(2.5))
    sec.different_first_page_header_footer = True
    fp = sec.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_field(fp, "PAGE")

    props = doc.core_properties
    props.title = content.TITLE_PAGE["title"]
    props.subject = content.TITLE_PAGE["course"]
    props.author = ""
    props.language = "pl-PL" if lang == "pl" else "en-GB"

    title_page(doc, content.TITLE_PAGE)

    fig_dir = HERE.parent / "build"
    for block in content.BODY:
        kind = block[0]
        if kind in ("h1", "h2", "h3"):
            doc.add_heading(block[1], level=int(kind[1]))
        elif kind == "p":
            add_rich(doc.add_paragraph(), block[1])
        elif kind == "bullets":
            for item in block[1]:
                add_rich(doc.add_paragraph(style="List Bullet"), item)
        elif kind == "numbered":
            for item in block[1]:
                add_rich(doc.add_paragraph(style="List Number"), item)
        elif kind == "refs":
            for i, item in enumerate(block[1], 1):
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Cm(0.9)
                p.paragraph_format.first_line_indent = Cm(-0.9)
                p.paragraph_format.line_spacing = 1.15
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                add_rich(p, f"[{i}]\t{item}", size=11)
        elif kind == "table":
            _, caption, headers, rows, widths = block
            doc.add_paragraph(caption, style="Caption").paragraph_format.keep_with_next = True
            add_table(doc, headers, rows, widths, bullets=True)
            doc.add_paragraph().paragraph_format.space_after = Pt(0)
        elif kind == "procon":
            _, caption, pros_title, pros, cons_title, cons = block
            doc.add_paragraph(caption, style="Caption").paragraph_format.keep_with_next = True
            t = add_table(doc, [pros_title, cons_title], [[pros, cons]], [8.0, 8.0], bullets=True)
            shade(t.rows[0].cells[0], "E2EFDA")
            shade(t.rows[0].cells[1], "FBE4D5")
            doc.add_paragraph().paragraph_format.space_after = Pt(0)
        elif kind == "figure":
            _, fname, caption = block
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.keep_with_next = True
            p.add_run().add_picture(str(fig_dir / fname.format(lang=lang)), width=Cm(16))
            doc.add_paragraph(caption, style="Caption")
        elif kind == "pagebreak":
            doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        else:
            raise ValueError(kind)

    doc.save(out_path)


if __name__ == "__main__":
    sys.path.insert(0, str(HERE))
    build(sys.argv[1], sys.argv[2])
